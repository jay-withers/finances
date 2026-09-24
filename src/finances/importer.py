"""Reading Finances_3.xlsx into a Document, once.

This runs from a laptop against the real blob and is then never needed again,
so it optimises for being *checkable* rather than for being reusable: every
total it produces is reported next to the spreadsheet's own, and anything it
declines to import is named rather than dropped.

`openpyxl` is a dev extra, not a runtime dependency — the deployed image has no
reason to carry a parser for a file format it never sees.

Three judgement calls, all reported by `reconcile()`:

1. **Pot ledgers are not reconstructed.** The `Savings` sheet holds month-end
   balances, not transactions: Holidays went 2700 -> 2900 in a month with a 450
   contribution, so 250 was spent, but the sheet does not say on what or when.
   Inventing a deposit and a withdrawal to explain each delta would fabricate
   history. The latest balance is imported as a single opening entry and the
   earlier months are kept as a note on it.

2. **The `Outgoings` savings block is not imported.** It duplicates the pot
   contributions from `Savings!B` and had drifted from them (Ted 25 against 50,
   Oven 52 against 55). The pots are the source; the block's total is reported
   for comparison and then discarded.

3. **Obviously-placeholder expiry dates become rolling contracts.** The two
   Lebara SIMs carry 2001-01-01 — twenty-five years past — because the sheet
   had no way to say "monthly, no end date".
"""

from __future__ import annotations

import calendar
import datetime
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from .model import (
    Document,
    Line,
    Pot,
    PotEntry,
    Renewal,
    Transfer,
    WealthAccount,
    WealthSnapshot,
)
from .money import format_money, parse_money

# A date this far back cannot be a real renewal date on a sheet maintained in
# 2026; it is the sheet's way of writing "no end date".
PLACEHOLDER_BEFORE = datetime.date(2010, 1, 1)

MONTH_COLUMNS = "CDEFGHIJKLMN"  # January..December on the Savings sheet


@dataclass
class Check:
    """One line of the reconciliation report."""

    label: str
    sheet: str
    imported: str
    ok: bool
    note: str = ""


@dataclass
class Result:
    document: Document
    checks: list[Check]
    skipped: list[str]


def _cell(sheet: Any, ref: str) -> Any:
    return sheet[ref].value


def _text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _as_date(value: Any) -> datetime.date | None:
    """A date from a cell openpyxl may hand back as datetime, date or number."""
    if value is None:
        return None
    if isinstance(value, datetime.datetime):
        return value.date()
    if isinstance(value, datetime.date):
        return value
    if isinstance(value, int | float):
        # Excel's epoch, with the 1900 leap-year bug already accounted for by
        # counting from 1899-12-30 rather than 1900-01-01.
        return datetime.date(1899, 12, 30) + datetime.timedelta(days=int(value))
    return None


def _end_of_month(year: int, month: int) -> datetime.date:
    return datetime.date(year, month, calendar.monthrange(year, month)[1])


def load_workbook(path: str | Path) -> Any:
    """The workbook with formulas resolved to their last-calculated values.

    `data_only=True` returns what Excel cached, which is why the file must have
    been saved by Excel rather than generated: a workbook written by a library
    has no cached values and every formula cell reads as None.
    """
    from openpyxl import load_workbook as _load

    return _load(path, data_only=True)


def import_workbook(path: str | Path, today: datetime.date | None = None) -> Result:
    """Read the whole spreadsheet. Returns the document *and* how to check it."""
    today = today or datetime.date.today()
    book = load_workbook(path)

    doc = Document()
    skipped: list[str] = []

    _import_income_and_outgoings(book["Outgoings"], doc, skipped)
    pot_ids, sheet_balances, last_month = _import_pots(book["Savings"], doc, today, skipped)
    _import_transfers(book["Monthly"], doc, pot_ids)
    _import_renewals(book["Expiry dates"], doc, skipped)
    _import_wealth(book["Wealth"], doc, skipped)

    checks = reconcile(book, doc, sheet_balances, last_month)
    return Result(document=doc, checks=checks, skipped=skipped)


def _import_income_and_outgoings(sheet: Any, doc: Document, skipped: list[str]) -> None:
    """Columns A:B are money in, D:E money out. G:H is deliberately skipped."""
    for row in range(2, 100):
        name = _text(_cell(sheet, f"A{row}"))
        if not name:
            break
        if name.lower() == "total":  # a formula cell, recomputed by calc.summary
            continue
        doc.income.append(Line(name=name, amount_pence=parse_money(_cell(sheet, f"B{row}"))))

    for row in range(2, 100):
        name = _text(_cell(sheet, f"D{row}"))
        amount = _cell(sheet, f"E{row}")
        if not name and amount is None:
            # A blank row mid-column: the sheet has one at D40 before its total.
            if _text(_cell(sheet, f"D{row + 1}")).lower() == "total":
                break
            continue
        if name.lower() == "total":
            continue
        if amount is None:
            # "Volleyball" has a name and no figure. Kept as an inactive line
            # rather than dropped: it is a commitment that currently costs
            # nothing, and deleting it loses the fact that it exists.
            doc.outgoings.append(Line(name=name, amount_pence=0, active=False, note="no amount"))
            skipped.append(f"outgoing {name!r} has no amount — imported as inactive")
            continue
        doc.outgoings.append(Line(name=name, amount_pence=parse_money(amount)))

    skipped.append(
        "Outgoings!G:H (the savings/spends block) not imported — it duplicates "
        "the pot contributions from Savings!B and had drifted from them"
    )


def _import_pots(
    sheet: Any, doc: Document, today: datetime.date, skipped: list[str]
) -> tuple[dict[str, str], dict[str, int], int]:
    """Rows 2..10 are pots; column B is the monthly amount, C..N month-end balances.

    Returns the name->id map, the imported balances, and the 1-based index of
    the last month that had any figures.
    """
    names: dict[str, str] = {}
    balances: dict[str, int] = {}

    rows: list[tuple[int, str]] = []
    for row in range(2, 40):
        name = _text(_cell(sheet, f"A{row}"))
        if not name:
            break
        if name.lower() == "total":
            break
        rows.append((row, name))

    # The last month with a figure against any pot. Read from the pot rows
    # rather than the sheet's own total row, which holds a formula that
    # evaluates to 0 for every empty month and so looks filled.
    last_month = 0
    for index, column in enumerate(MONTH_COLUMNS, start=1):
        if any(_cell(sheet, f"{column}{row}") is not None for row, _ in rows):
            last_month = index

    # The sheet carries no year. It is maintained in the current one, and the
    # last filled month is in the past, so: this year.
    year = today.year
    as_of = _end_of_month(year, last_month) if last_month else today

    for row, name in rows:
        monthly = parse_money(_cell(sheet, f"B{row}"))
        pot = Pot(
            name=name,
            monthly_pence=monthly,
            # A pot with no monthly contribution is the general savings pot,
            # which the sheet excluded from its own 925 total by summing B2:B9
            # rather than B2:B10.
            counts_toward_payday=monthly > 0,
        )
        doc.pots.append(pot)
        names[name] = pot.id

        if not last_month:
            continue
        balance = parse_money(_cell(sheet, f"{MONTH_COLUMNS[last_month - 1]}{row}"))
        balances[name] = balance

        history = []
        for index in range(last_month):
            value = _cell(sheet, f"{MONTH_COLUMNS[index]}{row}")
            if value is not None:
                month_name = calendar.month_abbr[index + 1]
                history.append(f"{month_name} {format_money(parse_money(value))}")

        doc.pot_entries.append(
            PotEntry(
                pot_id=pot.id,
                on=as_of,
                amount_pence=balance,
                kind="opening",
                note="Imported from the spreadsheet. Earlier months, which the "
                "sheet recorded as balances rather than movements: " + ", ".join(history),
            )
        )

    skipped.append(
        "Savings!I42 (a stray 2875 with no row label) not imported — "
        "it belongs to no pot and no column heading"
    )
    return names, balances, last_month


def _import_transfers(sheet: Any, doc: Document, pot_ids: dict[str, str]) -> None:
    """The payday routine: an ordered list, the last step being the remainder."""
    for row in range(1, 20):
        order = _cell(sheet, f"A{row}")
        label = _text(_cell(sheet, f"C{row}"))
        if order is None and not label:
            break
        raw = _cell(sheet, f"B{row}")
        note = _text(_cell(sheet, f"D{row}"))

        if isinstance(raw, str) and raw.strip().lower() == "remainder":
            doc.transfers.append(Transfer(label=label, amount_pence=None, note=note))
            continue

        amount = parse_money(raw)
        # The step whose figure came from Savings!B11. Recognised by value
        # rather than by reading the formula, because data_only=True gives the
        # cached result and not the expression that produced it.
        tracks = amount == sum(p.monthly_pence for p in doc.pots if p.counts_toward_payday)
        doc.transfers.append(
            Transfer(
                label=label,
                amount_pence=None if tracks else amount,
                note=note,
                tracks_pots=tracks,
            )
        )


def _import_renewals(sheet: Any, doc: Document, skipped: list[str]) -> None:
    for row in range(2, 100):
        kind = _text(_cell(sheet, f"A{row}"))
        if not kind:
            break
        expires = _as_date(_cell(sheet, f"C{row}"))
        company = _text(_cell(sheet, f"B{row}"))
        rolling = expires is not None and expires < PLACEHOLDER_BEFORE
        if rolling:
            skipped.append(
                f"renewal {kind!r} had a placeholder date of {expires:%Y-%m-%d} — "
                "imported as a rolling contract with no expiry"
            )
            expires = None
        doc.renewals.append(
            Renewal(
                kind=kind,
                company="" if company.upper() == "N/A" else company,
                expires_on=expires,
                rolling=rolling,
                comment=_text(_cell(sheet, f"D{row}")),
            )
        )


def _import_wealth(sheet: Any, doc: Document, skipped: list[str]) -> None:
    for row in range(2, 40):
        company = _text(_cell(sheet, f"A{row}"))
        if not company:
            break
        if company.lower() == "total":
            continue
        as_of = _as_date(_cell(sheet, f"B{row}"))
        if as_of is None:
            skipped.append(f"wealth account {company!r} has no 'as of' date — no snapshot imported")
            continue

        account = WealthAccount(company=company, notes=_text(_cell(sheet, f"G{row}")))
        doc.wealth_accounts.append(account)

        current = _cell(sheet, f"C{row}")
        projection = _cell(sheet, f"E{row}")
        growth = _cell(sheet, f"F{row}")
        doc.wealth_snapshots.append(
            WealthSnapshot(
                account_id=account.id,
                as_of=as_of,
                current_pence=parse_money(current) if current is not None else None,
                yearly_projection_pence=(
                    parse_money(projection) if projection is not None else None
                ),
                year_growth=float(growth) if growth is not None else None,
            )
        )


# --- checking the result ------------------------------------------------------


def reconcile(
    book: Any, doc: Document, sheet_balances: dict[str, int], last_month: int
) -> list[Check]:
    """Every imported total beside the spreadsheet's own.

    Two of these are *expected* to differ, and say so. The rest must match, and
    a mismatch means the import is wrong rather than the sheet being stale.
    """
    from . import calc

    checks: list[Check] = []
    outgoings = book["Outgoings"]
    savings = book["Savings"]

    totals = calc.summary(doc)

    sheet_income = parse_money(_cell(outgoings, "B4"))
    checks.append(
        Check(
            "Money in",
            format_money(sheet_income),
            format_money(totals.income),
            sheet_income == totals.income,
        )
    )

    sheet_out = parse_money(_cell(outgoings, "E41"))
    checks.append(
        Check(
            "Money out",
            format_money(sheet_out),
            format_money(totals.outgoings),
            sheet_out == totals.outgoings,
            # The sheet's own cell shows 2722.5299999999997: Excel summed 38
            # floats. Integer pence is why this one lands exactly.
            note="the sheet's cell carries float drift in its last places",
        )
    )

    sheet_block = parse_money(_cell(outgoings, "H13"))
    derived = totals.savings_and_spends
    checks.append(
        Check(
            "Savings / spends",
            format_money(sheet_block),
            format_money(derived),
            # Expected to differ — this is the drift the import exists to fix.
            True,
            note=(
                "EXPECTED DIFFERENCE. The sheet's block had drifted from "
                "Savings!B by "
                f"{format_money(abs(derived - sheet_block))}; the pot amounts win."
            )
            if derived != sheet_block
            else "",
        )
    )

    sheet_monthly = parse_money(_cell(savings, "B11"))
    imported_monthly = calc.pots_monthly_total(doc)
    checks.append(
        Check(
            "Pot contributions / month",
            format_money(sheet_monthly),
            format_money(imported_monthly),
            sheet_monthly == imported_monthly,
        )
    )

    if last_month:
        column = MONTH_COLUMNS[last_month - 1]
        sheet_total = parse_money(_cell(savings, f"{column}11"))
        imported_total = calc.pots_total(doc)
        checks.append(
            Check(
                f"Pot balances ({calendar.month_name[last_month]})",
                format_money(sheet_total),
                format_money(imported_total),
                sheet_total == imported_total,
            )
        )

    sheet_wealth = parse_money(_cell(book["Wealth"], "C6"))
    checks.append(
        Check(
            "Wealth total",
            format_money(sheet_wealth),
            format_money(calc.wealth_total(doc)),
            sheet_wealth == calc.wealth_total(doc),
        )
    )

    sheet_projection = parse_money(_cell(book["Wealth"], "E6"))
    checks.append(
        Check(
            "Yearly projection",
            format_money(sheet_projection),
            format_money(calc.projection_total(doc)),
            sheet_projection == calc.projection_total(doc),
        )
    )

    # Counted by walking the column rather than trusting `max_row`, which the
    # sheet's autofilter range inflates past the last row that holds anything.
    sheet_renewals = sum(1 for cell in book["Expiry dates"]["A"][1:] if _text(cell.value))
    checks.append(
        Check(
            "Renewals",
            str(sheet_renewals),
            str(len(doc.renewals)),
            len(doc.renewals) == sheet_renewals,
        )
    )

    # Spare is not compared against the sheet's own H15: that cell is computed
    # from the drifted block above, so it is wrong by the same amount and
    # matching it would mean reproducing the bug.
    checks.append(
        Check(
            "Spare",
            format_money(parse_money(_cell(outgoings, "H15"))),
            format_money(totals.spare),
            True,
            note="EXPECTED DIFFERENCE. Follows from the savings/spends line above.",
        )
    )

    return checks
